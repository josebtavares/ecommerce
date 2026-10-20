"""
Exportação dos relatórios (utils/relatorios_motor.py) em CSV, Excel e PDF — sempre a partir dos dados do
backend, com todas as linhas (não só as que o ecrã mostra). Os totais são os mesmos nos três formatos.
"""
import csv
import io
from datetime import date, datetime
from decimal import Decimal

from .horarios import fuso

MAX_LINHAS_PDF = 5000
AVISO_FISCAL = 'Documento de apoio à contabilidade. Não é um documento fiscal.'
MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']


# ── formatação ───────────────────────────────────────────────────────

def dinheiro(v):
    """1250.5 → "1.250$50" (escudo cabo-verdiano)."""
    if v is None or v == '':
        return ''
    n = Decimal(str(v)).quantize(Decimal('0.01'))
    inteiro, dec = f'{abs(n):.2f}'.split('.')
    inteiro = f'{int(inteiro):,}'.replace(',', '.')
    return f'{"-" if n < 0 else ""}{inteiro}${dec}'


def _local(v):
    return v.astimezone(fuso()) if isinstance(v, datetime) and v.tzinfo else v


def _data(v, tipo):
    if not v:
        return ''
    if isinstance(v, str):
        v = date.fromisoformat(v[:10])
    v = _local(v)
    if tipo == 'mes':
        return f'{MESES[v.month - 1]}/{v.year}'
    if tipo == 'datahora' and isinstance(v, datetime):
        return v.strftime('%d/%m/%Y %H:%M')
    return v.strftime('%d/%m/%Y')


def texto(v, tipo):
    """O valor como aparece no PDF e no ecrã."""
    if v is None:
        return ''
    if tipo == 'dinheiro':
        return dinheiro(v)
    if tipo in ('data', 'datahora', 'mes'):
        return _data(v, tipo)
    if tipo == 'percent':
        return f'{Decimal(str(v)).normalize():f}'.replace('.', ',') + '%'
    if tipo == 'id':                       # nº de encomenda, fatura… nunca com separador de milhares
        return str(v)
    if tipo == 'numero':
        if isinstance(v, (float, Decimal)) and v != int(v):
            return f'{v:.1f}'.replace('.', ',')
        return f'{int(v):,}'.replace(',', '.') if abs(int(v)) >= 10000 else str(int(v))
    return str(v)


def json_valor(v, tipo):
    """O valor para a API (o ecrã formata)."""
    if v is None:
        return None
    if isinstance(v, Decimal):
        return str(v.quantize(Decimal('0.01'))) if tipo in ('dinheiro', 'percent') else (int(v) if v == int(v) else float(v))
    if isinstance(v, datetime):
        return _local(v).isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return v


def _cabecalho(meta):
    """Linhas de identificação: emissor, período, filtros, quando e por quem."""
    out = [(meta['emissor'], meta.get('nif') and f'NIF {meta["nif"]}' or ''),
           ('Relatório', meta['titulo']),
           ('Situação em', _local(meta['gerado_em']).strftime('%d/%m/%Y %H:%M')) if meta.get('sem_periodo') else
           ('Período', f'{_data(meta["inicio"], "data")} a {_data(meta["fim"], "data")}')]
    out += [(k, v) for k, v in meta.get('filtros', [])]
    out += [('Gerado em', _local(meta['gerado_em']).strftime('%d/%m/%Y %H:%M')), ('Gerado por', meta['gerado_por'])]
    return out


# ── injecção de fórmulas ─────────────────────────────────────────────
# Nomes de clientes, produtos ou pesquisas vêm de quem usa o site: "=HYPERLINK(...)" ou "+cmd" num CSV/Excel
# seria executado pela folha de cálculo de quem o abre. No CSV leva um apóstrofo à frente; no Excel fica texto.
PERIGOSOS = ('=', '+', '-', '@', '\t', '\r')


def _texto_seguro(v):
    return "'" + v if isinstance(v, str) and v.startswith(PERIGOSOS) else v


# ── CSV ──────────────────────────────────────────────────────────────

def _csv_valor(v, tipo):
    """CSV para o Excel em português: ; como separador, vírgula decimal, datas dd/mm/aaaa."""
    if v is None:
        return ''
    if tipo in ('dinheiro', 'percent'):
        return f'{Decimal(str(v)).quantize(Decimal("0.01"))}'.replace('.', ',')
    if tipo in ('numero', 'id'):
        return str(v).replace('.', ',')
    return _texto_seguro(texto(v, tipo))


def gerar_csv(meta, res):
    buf = io.StringIO()
    buf.write('﻿')                                     # BOM: o Excel abre com os acentos certos
    w = csv.writer(buf, delimiter=';', lineterminator='\r\n')
    for k, v in _cabecalho(meta):
        w.writerow([_texto_seguro(k), _texto_seguro(v)])
    w.writerow([])
    cols = res['colunas']
    w.writerow([c['label'] for c in cols])
    for l in res['linhas']:
        w.writerow([_csv_valor(l.get(c['chave']), c['tipo']) for c in cols])
    if res.get('totais'):
        w.writerow([('Total' if i == 0 else _csv_valor(res['totais'].get(c['chave']), c['tipo']) if c['chave'] in res['totais'] else '')
                    for i, c in enumerate(cols)])
    return buf.getvalue().encode('utf-8')


# ── Excel ────────────────────────────────────────────────────────────

def gerar_xlsx(meta, res):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    negrito = Font(bold=True)
    fundo = PatternFill('solid', fgColor='EEEEEE')
    fmt = {'dinheiro': '#,##0.00', 'percent': '0.00"%"', 'data': 'DD/MM/YYYY', 'datahora': 'DD/MM/YYYY HH:MM',
           'mes': 'MMM/YYYY', 'numero': 'General'}

    def celula(ws, r, c, v, tipo):
        if v is None or v == '':
            return ws.cell(r, c, None)
        if tipo in ('dinheiro', 'percent', 'numero'):
            v = float(v)
        elif tipo == 'id':
            v = int(v)
        elif tipo in ('data', 'mes'):
            v = date.fromisoformat(v[:10]) if isinstance(v, str) else _local(v)
            if isinstance(v, datetime):
                v = v.date()
        elif tipo == 'datahora':
            v = _local(v).replace(tzinfo=None) if isinstance(v, datetime) else v
        cel = ws.cell(r, c, v)
        if isinstance(v, str):
            cel.data_type = 's'                 # nunca fórmula (ver PERIGOSOS)
        if tipo in fmt:
            cel.number_format = fmt[tipo]
        return cel

    # folha 1: identificação + resumo + notas
    ws = wb.active
    ws.title = 'Resumo'
    r = 1
    for k, v in _cabecalho(meta):
        celula(ws, r, 1, k, 'texto').font = negrito
        celula(ws, r, 2, v, 'texto')
        r += 1
    r += 1
    ws.cell(r, 1, 'Resumo').font = Font(bold=True, size=12)
    r += 1
    for item in res['resumo']:
        celula(ws, r, 1, item['label'], 'texto')
        celula(ws, r, 2, item['valor'], item['tipo'])
        r += 1
    r += 1
    for n in res.get('notas', []) + [AVISO_FISCAL]:
        ws.cell(r, 1, n).alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 15 * max(1, len(n) // 90 + 1)
        r += 1
    ws.column_dimensions['A'].width = 38
    ws.column_dimensions['B'].width = 30

    # folha 2: os dados, com filtro automático e cabeçalho fixo
    wd = wb.create_sheet('Dados')
    cols = res['colunas']
    for c, col in enumerate(cols, 1):
        cel = wd.cell(1, c, col['label'] + (' (CVE)' if col['tipo'] == 'dinheiro' else ''))
        cel.font = negrito
        cel.fill = fundo
    for i, l in enumerate(res['linhas'], 2):
        for c, col in enumerate(cols, 1):
            celula(wd, i, c, l.get(col['chave']), col['tipo'])
    n = len(res['linhas']) + 1
    if res.get('totais'):
        t = n + 1
        wd.cell(t, 1, 'Total').font = negrito
        for c, col in enumerate(cols, 1):
            if col['chave'] in res['totais'] and c > 1:
                celula(wd, t, c, res['totais'][col['chave']], col['tipo']).font = negrito
    wd.freeze_panes = 'A2'
    if n > 1:
        wd.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{n}'
    for c, col in enumerate(cols, 1):
        larg = max([len(col['label']) + 6] + [len(texto(l.get(col['chave']), col['tipo'])) for l in res['linhas'][:300]])
        wd.column_dimensions[get_column_letter(c)].width = min(max(10, larg + 2), 45)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────

def _pdf_txt(s):
    # as fontes base do PDF (Helvetica) não têm o sinal menos tipográfico
    return str(s).replace('−', '-').replace('…', '...')


def gerar_pdf(meta, res):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    cols = res['colunas']
    deitado = len(cols) > 7
    pagina = landscape(A4) if deitado else A4
    buf = io.BytesIO()
    titulo_doc = f'{meta["titulo"]} — {meta["emissor"]}'

    class Numerada(rl_canvas.Canvas):
        """Rodapé "Página X de Y" (precisa de saber o total no fim)."""
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._paginas = []

        def showPage(self):
            self._paginas.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._paginas)
            for estado in self._paginas:
                self.__dict__.update(estado)
                self.setFont('Helvetica', 7)
                self.setFillColor(colors.HexColor('#666666'))
                self.drawString(15 * mm, 8 * mm, _pdf_txt(f'{meta["emissor"]} · {meta["titulo"]} · {AVISO_FISCAL}'))
                self.drawRightString(pagina[0] - 15 * mm, 8 * mm, f'Página {self._pageNumber} de {total}')
                super().showPage()
            super().save()

    doc = SimpleDocTemplate(buf, pagesize=pagina, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=14 * mm,
                            bottomMargin=16 * mm, title=_pdf_txt(titulo_doc), author=_pdf_txt(meta['emissor']))
    est_titulo = ParagraphStyle('t', fontName='Helvetica-Bold', fontSize=15, leading=19, spaceAfter=2)
    est_sub = ParagraphStyle('s', fontName='Helvetica', fontSize=8.5, leading=11, textColor=colors.HexColor('#444444'))
    est_cel = ParagraphStyle('c', fontName='Helvetica', fontSize=7, leading=8.5)
    est_cab = ParagraphStyle('h', fontName='Helvetica-Bold', fontSize=7, leading=8.5)
    est_nota = ParagraphStyle('n', fontName='Helvetica-Oblique', fontSize=7.5, leading=9.5, textColor=colors.HexColor('#555555'))
    esc = lambda s: _pdf_txt(s).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    hist = [Paragraph(esc(meta['titulo']), est_titulo)]
    for k, v in _cabecalho(meta):
        if k == 'Relatório':
            continue                    # já é o título
        hist.append(Paragraph(f'<b>{esc(k)}</b>  {esc(v)}', est_sub))
    hist.append(Spacer(1, 5 * mm))

    # resumo em grelha (até 4 por linha)
    if res['resumo']:
        por_linha = 4
        cel = [[Paragraph(f'<font size=7 color="#666666">{esc(i["label"])}</font><br/><b>{esc(texto(i["valor"], i["tipo"]) or "-")}</b>',
                          ParagraphStyle('r', fontName='Helvetica', fontSize=10, leading=13)) for i in res['resumo'][k:k + por_linha]]
               for k in range(0, len(res['resumo']), por_linha)]
        for linha in cel:
            linha += [''] * (por_linha - len(linha))
        larg = (pagina[0] - 30 * mm) / por_linha
        t = Table(cel, colWidths=[larg] * por_linha)
        t.setStyle(TableStyle([('BOX', (0, 0), (-1, -1), 0.4, colors.HexColor('#cccccc')),
                               ('INNERGRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#e2e2e2')),
                               ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 4),
                               ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
        hist += [t, Spacer(1, 5 * mm)]

    linhas = res['linhas'][:MAX_LINHAS_PDF]
    direita = {i for i, c in enumerate(cols) if c['tipo'] in ('dinheiro', 'numero', 'percent', 'id')}
    dados = [[Paragraph(esc(c['label']), est_cab) for c in cols]]
    for l in linhas:
        dados.append([esc(texto(l.get(c['chave']), c['tipo'])) if i in direita or c['tipo'] in ('data', 'datahora', 'mes')
                      else Paragraph(esc(texto(l.get(c['chave']), c['tipo'])), est_cel) for i, c in enumerate(cols)])
    if res.get('totais') and linhas:
        dados.append([('Total' if i == 0 else esc(texto(res['totais'][c['chave']], c['tipo'])) if c['chave'] in res['totais'] else '')
                      for i, c in enumerate(cols)])
    if not linhas:
        dados.append([Paragraph('Sem resultados para estes filtros.', est_cel)] + [''] * (len(cols) - 1))

    # larguras: números e datas cabem sempre inteiros (medidos); as colunas de texto ficam com o que sobra
    from reportlab.pdfbase.pdfmetrics import stringWidth
    util = pagina[0] - 30 * mm
    fixas = {}
    for i, c in enumerate(cols):
        if c['tipo'] in ('dinheiro', 'numero', 'percent', 'id', 'data', 'datahora', 'mes'):
            amostra = [texto(l.get(c['chave']), c['tipo']) for l in linhas[:300]]
            if res.get('totais') and c['chave'] in res['totais']:
                amostra.append(texto(res['totais'][c['chave']], c['tipo']))
            conteudo = max([stringWidth(_pdf_txt(x), 'Helvetica-Bold', 7) for x in amostra] + [0])
            cab = max(stringWidth(_pdf_txt(p_), 'Helvetica-Bold', 7) for p_ in c['label'].split())
            fixas[i] = max(conteudo, cab) + 8
    resto = util - sum(fixas.values())
    n_texto = len(cols) - len(fixas)
    if n_texto:
        larg_texto = max(resto / n_texto, 18 * mm)
        larguras = [fixas.get(i, larg_texto) for i in range(len(cols))]
    else:
        larguras = [fixas[i] for i in range(len(cols))]
    escala = util / sum(larguras)
    if escala < 1 or not n_texto:
        larguras = [w * escala for w in larguras]

    t = Table(dados, colWidths=larguras, repeatRows=1, hAlign='LEFT')
    estilo = [('FONTNAME', (0, 0), (-1, -1), 'Helvetica'), ('FONTSIZE', (0, 0), (-1, -1), 7),
              ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#eeeeee')),
              ('LINEBELOW', (0, 0), (-1, 0), 0.6, colors.HexColor('#999999')),
              ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('TOPPADDING', (0, 0), (-1, -1), 2.5),
              ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5), ('LEFTPADDING', (0, 0), (-1, -1), 3), ('RIGHTPADDING', (0, 0), (-1, -1), 3),
              ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f7f7f7')])]
    for i in direita:
        estilo.append(('ALIGN', (i, 1), (i, -1), 'RIGHT'))
    if res.get('totais') and linhas:
        estilo += [('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'), ('LINEABOVE', (0, -1), (-1, -1), 0.8, colors.black),
                   ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#eeeeee'))]
    t.setStyle(TableStyle(estilo))
    hist.append(t)

    hist.append(Spacer(1, 4 * mm))
    if len(res['linhas']) > MAX_LINHAS_PDF:
        hist.append(Paragraph(esc(f'O PDF mostra as primeiras {MAX_LINHAS_PDF} linhas de {len(res["linhas"])} '
                                  '(os totais são de todas). Para a lista completa usa o Excel ou o CSV.'), est_nota))
    for n in res.get('notas', []):
        hist.append(Paragraph(esc(n), est_nota))

    doc.build(hist, canvasmaker=Numerada)
    return buf.getvalue()


FORMATOS = {
    'csv':  (gerar_csv, 'text/csv; charset=utf-8', 'csv'),
    'xlsx': (gerar_xlsx, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'xlsx'),
    'pdf':  (gerar_pdf, 'application/pdf', 'pdf'),
}
